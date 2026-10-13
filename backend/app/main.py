import json
import io
import csv
from datetime import datetime, timezone
from fastapi import FastAPI, Depends, HTTPException, Header, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response, JSONResponse
from sqlalchemy.orm import Session
from sqlalchemy import select, text
from .config import settings
from .database import Base, engine, get_db
from .models import User, Assessment, PortfolioEntry
from .schemas import RegisterInput, LoginInput, TokenResponse, UserPublic, StartupInput, AssessmentResponse, CompareInput
from .security import hash_password, verify_password, create_access_token, decode_access_token
from .assessment import assess_startup
from .reporting import build_pdf
from .pitch_deck import analyze_pitch_text
from .llm_review import review_with_llm
from .llm_extract import extract_with_llm
from .analytics import build_analytics
from .investment_memo import build_memo, build_investment_memo_pdf
from .decision_support import build_decision_support
from .data_management import validate_assessment_backup
from pypdf import PdfReader

app = FastAPI(title="AI Venture Assessment System", version="1.0.0",
              description="Educational MVP. Synthetic model predictions are not validated startup-success forecasts.")
app.add_middleware(CORSMiddleware, allow_origins=list(dict.fromkeys([settings.frontend_origin, "http://127.0.0.1:5173"])),
                   allow_credentials=True, allow_methods=["GET", "POST", "PUT", "PATCH", "DELETE", "OPTIONS"],
                   allow_headers=["Authorization", "Content-Type"]) 

@app.middleware("http")
async def add_security_headers(request, call_next):
    response = await call_next(request)
    response.headers.setdefault("X-Content-Type-Options", "nosniff")
    response.headers.setdefault("X-Frame-Options", "DENY")
    response.headers.setdefault("Referrer-Policy", "strict-origin-when-cross-origin")
    response.headers.setdefault("Permissions-Policy", "camera=(), microphone=(), geolocation=()")
    response.headers.setdefault("Cache-Control", "no-store" if request.url.path.startswith("/api/auth") else "no-cache")
    if settings.environment.lower() == "production":
        response.headers.setdefault("Strict-Transport-Security", "max-age=31536000; includeSubDomains")
    return response

@app.on_event("startup")
def startup():
    if settings.environment.lower() == "production":
        if len(settings.jwt_secret) < 32 or settings.jwt_secret == "local-development-only-change-this-secret":
            raise RuntimeError("Production requires a strong random JWT_SECRET of at least 32 characters.")
        if settings.frontend_origin.strip() == "*":
            raise RuntimeError("Production FRONTEND_ORIGIN must be a specific trusted origin.")
    Base.metadata.create_all(bind=engine)

def current_user(authorization: str | None = Header(default=None), db: Session = Depends(get_db)) -> User:
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(status_code=401, detail="Sign in to continue.")
    user_id = decode_access_token(authorization.split(" ",1)[1].strip())
    if user_id is None:
        raise HTTPException(status_code=401, detail="Session is invalid or expired. Please sign in again.")
    user = db.get(User, user_id)
    if not user:
        raise HTTPException(status_code=401, detail="Account not found.")
    return user


@app.post("/api/pitch-deck/analyze")
async def analyze_pitch_deck(file: UploadFile = File(...), user: User = Depends(current_user)):
    """Extract and review a PDF pitch deck in memory; uploaded bytes are not persisted."""
    filename = (file.filename or "").lower()
    if not filename.endswith(".pdf"):
        raise HTTPException(status_code=415, detail="Please upload a PDF file.")
    content = await file.read(10 * 1024 * 1024 + 1)
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="PDF must be 10 MB or smaller.")
    if not content.startswith(b"%PDF-"):
        raise HTTPException(status_code=400, detail="The uploaded file does not appear to be a valid PDF.")
    try:
        reader = PdfReader(io.BytesIO(content), strict=False)
        if len(reader.pages) == 0:
            raise HTTPException(status_code=400, detail="The PDF has no readable pages.")
        if len(reader.pages) > 50:
            raise HTTPException(status_code=413, detail="PDF must contain 50 pages or fewer.")
        extracted = "\n".join((page.extract_text() or "") for page in reader.pages)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read this PDF. Try exporting it again as a text-based PDF.")
    if len(extracted.strip()) < 40:
        raise HTTPException(status_code=422, detail="Very little selectable text was found. Scanned/image-only PDFs are not supported yet.")
    return analyze_pitch_text(extracted, len(reader.pages))

@app.post("/api/pitch-deck/ai-review")
async def ai_review_pitch_deck(file: UploadFile = File(...), user: User = Depends(current_user)):
    """Optional LLM review. The PDF is extracted in memory and is not persisted."""
    if not settings.llm_api_key.strip():
        raise HTTPException(status_code=503, detail="AI review is not configured. Add LLM_API_KEY to backend/.env and restart the backend.")
    filename = (file.filename or "").lower()
    if not filename.endswith(".pdf"):
        raise HTTPException(status_code=415, detail="Please upload a PDF file.")
    content = await file.read(10 * 1024 * 1024 + 1)
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="PDF must be 10 MB or smaller.")
    if not content.startswith(b"%PDF-"):
        raise HTTPException(status_code=400, detail="The uploaded file does not appear to be a valid PDF.")
    try:
        reader = PdfReader(io.BytesIO(content), strict=False)
        if not reader.pages:
            raise HTTPException(status_code=400, detail="The PDF has no readable pages.")
        if len(reader.pages) > 50:
            raise HTTPException(status_code=413, detail="PDF must contain 50 pages or fewer.")
        extracted = "\n".join((page.extract_text() or "") for page in reader.pages)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read this PDF. Try exporting it again as a text-based PDF.")
    if len(extracted.strip()) < 40:
        raise HTTPException(status_code=422, detail="Very little selectable text was found. Scanned/image-only PDFs are not supported yet.")
    try:
        return await review_with_llm(extracted, len(reader.pages), settings.llm_api_key,
                                     settings.llm_base_url, settings.llm_model)
    except ValueError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    except RuntimeError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc


@app.post("/api/pitch-deck/extract-data")
async def extract_pitch_deck_data(file: UploadFile = File(...), user: User = Depends(current_user)):
    """Extract draft structured fields from a PDF in memory; never persists the PDF."""
    if not settings.llm_api_key.strip():
        raise HTTPException(status_code=503, detail="Data extraction AI is not configured. Add LLM_API_KEY to backend/.env and restart the backend.")
    filename = (file.filename or "").lower()
    if not filename.endswith(".pdf"):
        raise HTTPException(status_code=415, detail="Please upload a PDF file.")
    content = await file.read(10 * 1024 * 1024 + 1)
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="PDF must be 10 MB or smaller.")
    if not content.startswith(b"%PDF-"):
        raise HTTPException(status_code=400, detail="The uploaded file does not appear to be a valid PDF.")
    try:
        reader = PdfReader(io.BytesIO(content), strict=False)
        if not reader.pages:
            raise HTTPException(status_code=400, detail="The PDF has no readable pages.")
        if len(reader.pages) > 50:
            raise HTTPException(status_code=413, detail="PDF must contain 50 pages or fewer.")
        extracted = "\n".join((page.extract_text() or "") for page in reader.pages)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read this PDF. Try exporting it again as a text-based PDF.")
    if len(extracted.strip()) < 40:
        raise HTTPException(status_code=422, detail="Very little selectable text was found. Scanned/image-only PDFs are not supported yet.")
    try:
        result = await extract_with_llm(extracted, len(reader.pages), settings.llm_api_key,
                                        settings.llm_base_url, settings.llm_model)
        result["page_count"] = len(reader.pages)
        return result
    except ValueError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    except RuntimeError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc


@app.get("/api/health")
def health():
    """Liveness/readiness check. Does not expose database URLs or secret values."""
    try:
        with engine.connect() as connection:
            connection.execute(text("SELECT 1"))
    except Exception:
        raise HTTPException(status_code=503, detail="Database is unavailable.")
    return {"status": "ok", "service": "AVAS", "database": "connected",
            "environment": settings.environment.lower(),
            "ai_review_configured": bool(settings.llm_api_key.strip())}

@app.post("/api/auth/register", response_model=TokenResponse)
def register(payload: RegisterInput, db: Session = Depends(get_db)):
    email = payload.email.lower()
    if db.scalar(select(User).where(User.email == email)):
        raise HTTPException(status_code=409, detail="An account with this email already exists.")
    user = User(full_name=payload.full_name, email=email, password_hash=hash_password(payload.password))
    db.add(user)
    db.commit()
    db.refresh(user)
    return {"access_token":create_access_token(user.id),"user":user}

@app.post("/api/auth/login", response_model=TokenResponse)
def login(payload: LoginInput, db: Session = Depends(get_db)):
    user = db.scalar(select(User).where(User.email == payload.email.lower()))
    if not user or not verify_password(payload.password, user.password_hash):
        raise HTTPException(status_code=401, detail="Incorrect email or password.")
    return {"access_token":create_access_token(user.id),"user":user}

@app.get("/api/auth/me", response_model=UserPublic)
def me(user: User = Depends(current_user)):
    return user

@app.post("/api/assess", response_model=AssessmentResponse)
def assess(payload: StartupInput, user: User = Depends(current_user), db: Session = Depends(get_db)):
    result = assess_startup(payload.model_dump())
    # Preserve a reviewed snapshot for later memo/report generation.
    result["input_snapshot"] = payload.model_dump()
    record = Assessment(owner_id=user.id, startup_name=result["startup_name"], industry=result["industry"],
                       readiness_score=result["readiness_score"], risk_level=result["risk_level"],
                       result_json=json.dumps(result))
    db.add(record)
    db.commit()
    db.refresh(record)
    result["id"] = record.id
    result["created_at"] = record.created_at.isoformat() if record.created_at else None
    return result

@app.get("/api/assessments")
def history(user: User = Depends(current_user), db: Session = Depends(get_db)):
    rows = db.scalars(select(Assessment).where(Assessment.owner_id == user.id).order_by(Assessment.id.desc()).limit(100)).all()
    return [{"id":r.id,"startup_name":r.startup_name,"industry":r.industry,
             "readiness_score":r.readiness_score,"risk_level":r.risk_level,
             "created_at":r.created_at.isoformat() if r.created_at else None} for r in rows]

@app.post("/api/compare")
def compare_ventures(payload: CompareInput, user: User = Depends(current_user), db: Session = Depends(get_db)):
    """Compare 2–4 saved assessments belonging to the signed-in user."""
    if len(set(payload.assessment_ids)) != len(payload.assessment_ids):
        raise HTTPException(status_code=422, detail="Choose each assessment only once.")
    rows = db.scalars(select(Assessment).where(
        Assessment.owner_id == user.id, Assessment.id.in_(payload.assessment_ids)
    )).all()
    by_id = {row.id: row for row in rows}
    if len(by_id) != len(payload.assessment_ids):
        raise HTTPException(status_code=404, detail="One or more assessments were not found in your account.")
    items = []
    for assessment_id in payload.assessment_ids:
        row = by_id[assessment_id]
        try:
            result = json.loads(row.result_json)
        except (TypeError, json.JSONDecodeError):
            result = {}
        explain = result.get("explainability") or {}
        financial = explain.get("financial_snapshot") or {}
        items.append({
            "id": row.id, "startup_name": row.startup_name, "industry": row.industry,
            "readiness_score": row.readiness_score, "risk_level": row.risk_level,
            "growth_potential_score": result.get("growth_potential_score", row.readiness_score),
            "risk_flags": result.get("risk_flags", []),
            "positive_signals": explain.get("positive_signals", []),
            "priority_gaps": explain.get("priority_gaps", []),
            "recommendations": result.get("recommendations", []),
            "runway_months": financial.get("runway_months"),
            "monthly_net_burn": financial.get("monthly_net_burn"),
            "created_at": row.created_at.isoformat() if row.created_at else None,
        })
    scores = [item["readiness_score"] for item in items]
    highest = max(items, key=lambda item: item["readiness_score"])
    lowest = min(items, key=lambda item: item["readiness_score"])
    return {
        "items": items,
        "summary": {
            "compared_count": len(items),
            "score_spread": max(scores) - min(scores),
            "highest_readiness_id": highest["id"],
            "lowest_readiness_id": lowest["id"],
        },
        "notice": "Comparison summarizes saved rubric-based assessments. A higher readiness score is not proof of investment quality or future success. Verify assumptions and compare financial units consistently."
    }

@app.get("/api/decision-support")
def investment_decision_support(user: User = Depends(current_user), db: Session = Depends(get_db)):
    """Portfolio-level risk screening, priority alerts, and same-name assessment trends."""
    assessments = db.scalars(select(Assessment).where(Assessment.owner_id == user.id)
                             .order_by(Assessment.created_at.asc(), Assessment.id.asc())).all()
    portfolio_entries = db.scalars(select(PortfolioEntry).where(PortfolioEntry.owner_id == user.id)).all()
    return build_decision_support(assessments, portfolio_entries)


@app.get("/api/analytics")
def venture_analytics(user: User = Depends(current_user), db: Session = Depends(get_db)):
    """Aggregated analytics over the signed-in user's saved assessments only."""
    rows = db.scalars(select(Assessment).where(Assessment.owner_id == user.id).order_by(Assessment.created_at.asc(), Assessment.id.asc())).all()
    return build_analytics(rows)

@app.get("/api/assessments/{assessment_id}", response_model=AssessmentResponse)
def get_assessment(assessment_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    row = db.scalar(select(Assessment).where(Assessment.id == assessment_id, Assessment.owner_id == user.id))
    if not row:
        raise HTTPException(status_code=404, detail="Assessment not found.")
    result = json.loads(row.result_json)
    # Backward compatibility for assessments saved before the explainability update.
    result.setdefault("growth_potential_score", result.get("readiness_score", 0))
    result.setdefault("growth_score_note", "Legacy saved assessment: growth indicator was not calculated when this record was created.")
    result.setdefault("explainability", {
        "method": "This assessment was saved before the explainability update.",
        "positive_signals": ["Run a new assessment to generate current positive signals."],
        "priority_gaps": ["Run a new assessment to generate current priority gaps."],
        "financial_snapshot": {"runway_months": 0, "monthly_net_burn": 0, "revenue_covers_burn": False, "runway_alert": "Not available for this older assessment"},
        "model_explanation": {"available": False, "method": "SHAP TreeExplainer (optional)", "note": "This older assessment was saved before SHAP explanations were added.", "top_positive": [], "top_negative": []}
    })
    result["explainability"].setdefault("model_explanation", {"available": False, "method": "SHAP TreeExplainer (optional)", "note": "This older assessment does not contain a SHAP explanation. Run a new assessment.", "top_positive": [], "top_negative": []})
    result["id"] = row.id
    result["created_at"] = row.created_at.isoformat() if row.created_at else None
    return result

@app.get("/api/assessments/{assessment_id}/investment-memo")
def investment_memo(assessment_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    """Build a structured, account-scoped investment memo from one saved assessment."""
    row = db.scalar(select(Assessment).where(Assessment.id == assessment_id, Assessment.owner_id == user.id))
    if not row:
        raise HTTPException(status_code=404, detail="Assessment not found.")
    try:
        result = json.loads(row.result_json)
    except (TypeError, json.JSONDecodeError):
        raise HTTPException(status_code=422, detail="Saved assessment data could not be read.")
    return build_memo(result, row.id, row.created_at.isoformat() if row.created_at else None)

@app.get("/api/assessments/{assessment_id}/investment-memo.pdf")
def investment_memo_pdf(assessment_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    """Download the structured memo as a PDF, scoped to the signed-in account."""
    row = db.scalar(select(Assessment).where(Assessment.id == assessment_id, Assessment.owner_id == user.id))
    if not row:
        raise HTTPException(status_code=404, detail="Assessment not found.")
    try:
        result = json.loads(row.result_json)
    except (TypeError, json.JSONDecodeError):
        raise HTTPException(status_code=422, detail="Saved assessment data could not be read.")
    memo = build_memo(result, row.id, row.created_at.isoformat() if row.created_at else None)
    pdf = build_investment_memo_pdf(memo)
    filename = f"avas_investment_memo_{row.id}.pdf"
    return Response(content=pdf, media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})

@app.get("/api/assessments/{assessment_id}/report.pdf")
def report(assessment_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    row = db.scalar(select(Assessment).where(Assessment.id == assessment_id, Assessment.owner_id == user.id))
    if not row:
        raise HTTPException(status_code=404, detail="Assessment not found.")
    result = json.loads(row.result_json)
    pdf = build_pdf(result)
    filename = f"avas_assessment_{row.id}.pdf"
    return Response(content=pdf, media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@app.get("/api/reports/assessments.csv")
def export_assessments_csv(user: User = Depends(current_user), db: Session = Depends(get_db)):
    """Export an account-scoped assessment summary as CSV."""
    rows = db.scalars(select(Assessment).where(Assessment.owner_id == user.id).order_by(Assessment.id.asc())).all()
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(["Assessment ID", "Startup", "Industry", "Readiness Score", "Risk Level", "Created At"])
    for row in rows:
        writer.writerow([row.id, row.startup_name, row.industry, row.readiness_score, row.risk_level,
                         row.created_at.isoformat() if row.created_at else ""])
    return Response(content="\ufeff" + stream.getvalue(), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": 'attachment; filename="avas_assessments.csv"'})


@app.get("/api/reports/assessments.xlsx")
def export_assessments_xlsx(user: User = Depends(current_user), db: Session = Depends(get_db)):
    """Export summary and financial snapshot worksheets in an Excel workbook."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise HTTPException(status_code=503, detail="Excel export dependency is missing. Install backend requirements and restart.") from exc
    rows = db.scalars(select(Assessment).where(Assessment.owner_id == user.id).order_by(Assessment.id.asc())).all()
    wb = Workbook()
    summary = wb.active
    summary.title = "Assessments"
    summary.append(["Assessment ID", "Startup", "Industry", "Readiness Score", "Risk Level", "Created At"])
    financial = wb.create_sheet("Financial snapshots")
    financial.append(["Assessment ID", "Startup", "Industry", "Monthly Revenue", "Monthly Burn", "Funding Needed", "Runway Months", "Currency note"])
    for row in rows:
        summary.append([row.id, row.startup_name, row.industry, row.readiness_score, row.risk_level,
                        row.created_at.isoformat() if row.created_at else ""])
        try:
            result = json.loads(row.result_json)
        except (TypeError, json.JSONDecodeError):
            result = {}
        snap = result.get("input_snapshot") or {}
        financial.append([row.id, row.startup_name, row.industry, snap.get("monthly_revenue"), snap.get("monthly_burn"),
                          snap.get("funding_needed"), snap.get("financial_runway_months"), "As entered by user; currency not standardized"])
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="26324A")
            cell.alignment = Alignment(wrap_text=True)
        for column in sheet.columns:
            width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 42)
            sheet.column_dimensions[get_column_letter(column[0].column)].width = width
    output = io.BytesIO()
    wb.save(output)
    return Response(content=output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="avas_assessments.xlsx"'})


@app.get("/api/reports/backup.json")
def export_assessments_backup(user: User = Depends(current_user), db: Session = Depends(get_db)):
    """Download portable JSON backup of assessment records (not account credentials)."""
    rows = db.scalars(select(Assessment).where(Assessment.owner_id == user.id).order_by(Assessment.id.asc())).all()
    items = []
    for row in rows:
        try:
            result = json.loads(row.result_json)
        except (TypeError, json.JSONDecodeError):
            result = {}
        items.append({"startup_name": row.startup_name, "industry": row.industry,
                      "readiness_score": row.readiness_score, "risk_level": row.risk_level,
                      "created_at": row.created_at.isoformat() if row.created_at else None, "result": result})
    payload = {"format": "avas-assessment-backup", "schema_version": 1,
               "exported_at": datetime.now(timezone.utc).isoformat(), "assessments": items}
    return JSONResponse(content=payload, headers={"Content-Disposition": 'attachment; filename="avas_assessment_backup.json"'})


@app.post("/api/reports/import.json", status_code=201)
def import_assessments_backup(payload: dict, user: User = Depends(current_user), db: Session = Depends(get_db)):
    """Import validated assessment records into the signed-in user's account; IDs/owners are never imported."""
    try:
        items = validate_assessment_backup(payload)
    except OverflowError as exc:
        raise HTTPException(status_code=413, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    prepared = [Assessment(owner_id=user.id, startup_name=item["startup_name"], industry=item["industry"],
                           readiness_score=item["readiness_score"], risk_level=item["risk_level"],
                           result_json=json.dumps(item["result"])) for item in items]
    db.add_all(prepared)
    db.commit()
    return {"imported_count": len(prepared), "message": f"Imported {len(prepared)} assessment record(s). Existing records were not changed."}


PORTFOLIO_STATUSES = {"Watching", "Due diligence", "Active", "On hold", "Passed"}

@app.get("/api/portfolio")
def get_portfolio(user: User = Depends(current_user), db: Session = Depends(get_db)):
    """Return portfolio records and their latest saved assessment snapshot for this account."""
    rows = db.execute(
        select(PortfolioEntry, Assessment)
        .join(Assessment, PortfolioEntry.assessment_id == Assessment.id)
        .where(PortfolioEntry.owner_id == user.id, Assessment.owner_id == user.id)
        .order_by(PortfolioEntry.updated_at.desc(), PortfolioEntry.id.desc())
    ).all()
    return [{
        "id": entry.id, "assessment_id": assessment.id,
        "startup_name": assessment.startup_name, "industry": assessment.industry,
        "readiness_score": assessment.readiness_score, "risk_level": assessment.risk_level,
        "assessment_created_at": assessment.created_at.isoformat() if assessment.created_at else None,
        "status": entry.status, "notes": entry.notes,
        "created_at": entry.created_at.isoformat() if entry.created_at else None,
        "updated_at": entry.updated_at.isoformat() if entry.updated_at else None,
    } for entry, assessment in rows]

@app.post("/api/portfolio", status_code=201)
def add_to_portfolio(payload: dict, user: User = Depends(current_user), db: Session = Depends(get_db)):
    assessment_id = payload.get("assessment_id")
    if not isinstance(assessment_id, int) or isinstance(assessment_id, bool):
        raise HTTPException(status_code=422, detail="assessment_id must be an integer.")
    status = payload.get("status", "Watching")
    notes = payload.get("notes", "")
    if not isinstance(status, str) or status not in PORTFOLIO_STATUSES:
        raise HTTPException(status_code=422, detail="Choose a valid portfolio status.")
    if not isinstance(notes, str) or len(notes) > 5000:
        raise HTTPException(status_code=422, detail="Notes must be text with at most 5,000 characters.")
    assessment = db.scalar(select(Assessment).where(Assessment.id == assessment_id, Assessment.owner_id == user.id))
    if not assessment:
        raise HTTPException(status_code=404, detail="Assessment not found in your account.")
    exists = db.scalar(select(PortfolioEntry).where(PortfolioEntry.owner_id == user.id, PortfolioEntry.assessment_id == assessment_id))
    if exists:
        raise HTTPException(status_code=409, detail="This assessment is already in your portfolio.")
    entry = PortfolioEntry(owner_id=user.id, assessment_id=assessment_id, status=status, notes=notes.strip())
    db.add(entry); db.commit(); db.refresh(entry)
    return {"id": entry.id, "assessment_id": assessment.id, "startup_name": assessment.startup_name,
            "industry": assessment.industry, "readiness_score": assessment.readiness_score,
            "risk_level": assessment.risk_level, "assessment_created_at": assessment.created_at.isoformat() if assessment.created_at else None,
            "status": entry.status, "notes": entry.notes,
            "created_at": entry.created_at.isoformat() if entry.created_at else None,
            "updated_at": entry.updated_at.isoformat() if entry.updated_at else None}

@app.patch("/api/portfolio/{entry_id}")
def update_portfolio_entry(entry_id: int, payload: dict, user: User = Depends(current_user), db: Session = Depends(get_db)):
    entry = db.scalar(select(PortfolioEntry).where(PortfolioEntry.id == entry_id, PortfolioEntry.owner_id == user.id))
    if not entry:
        raise HTTPException(status_code=404, detail="Portfolio entry not found.")
    if "status" in payload:
        if not isinstance(payload["status"], str) or payload["status"] not in PORTFOLIO_STATUSES:
            raise HTTPException(status_code=422, detail="Choose a valid portfolio status.")
        entry.status = payload["status"]
    if "notes" in payload:
        if not isinstance(payload["notes"], str) or len(payload["notes"] ) > 5000:
            raise HTTPException(status_code=422, detail="Notes must be text with at most 5,000 characters.")
        entry.notes = payload["notes"].strip()
    db.commit(); db.refresh(entry)
    return {"id": entry.id, "status": entry.status, "notes": entry.notes,
            "updated_at": entry.updated_at.isoformat() if entry.updated_at else None}

@app.delete("/api/portfolio/{entry_id}", status_code=204)
def remove_from_portfolio(entry_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    entry = db.scalar(select(PortfolioEntry).where(PortfolioEntry.id == entry_id, PortfolioEntry.owner_id == user.id))
    if not entry:
        raise HTTPException(status_code=404, detail="Portfolio entry not found.")
    db.delete(entry); db.commit()
    return Response(status_code=204)
